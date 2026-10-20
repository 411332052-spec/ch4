"""Excel 紀錄：一直往下加新的一列，並算出當日已攝取熱量。

四段防護裡的「防護 4」在這個檔案。
"""
from __future__ import annotations

import asyncio
import logging
import os
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell

from .config import EXCEL_RETRIES, EXCEL_RETRY_DELAY_SECONDS

log = logging.getLogger(__name__)

HEADERS = [
    "日期時間",
    "上傳者",
    "食物敘述",
    "成份",
    "重量(公克)",
    "熱量(大卡)",
    "當日已攝取熱量(大卡)",
    "照片存放位置",
]
COL_WHEN, COL_KCAL = 1, 6  # 欄位編號（從 1 算起）
COLUMN_WIDTHS = [20, 16, 24, 36, 12, 12, 22, 60]


class LedgerBusyError(Exception):
    """Excel 一直寫不進去（通常是檔案正被打開著）。"""


@dataclass(frozen=True)
class MealRecord:
    when: datetime  # 上傳時間（本機時間）
    uploader: str
    description: str
    ingredients: str
    weight_g: int
    calories_kcal: int
    photo_path: str


def _write_text(cell: Cell, value: object) -> None:
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    cell.value = text  # 沒有防護時：原樣寫入，以 = 開頭的文字會被 Excel 當成公式
    # ==== 防護 4：寫進 Excel 的文字防公式注入　開始 ====
    text = " ".join(text.split("\n")).replace("\r", " ")  # 一格只放一行
    if text[:1] in ("=", "+", "-", "@", "\t"):
        text = "'" + text  # 前面加單引號，Excel 就不會把它當公式
    cell.value = text
    cell.data_type = "s"  # 明確指定為文字，不讓 Excel 自己猜
    # ==== 防護 4 結束 ====


def _as_datetime(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _daily_total(ws, day: date) -> float:
    """當天（本機日期）所有列的熱量加總。整個頻道一起算，不分人。"""
    total = 0.0
    for row in ws.iter_rows(min_row=2, values_only=True):
        when = _as_datetime(row[COL_WHEN - 1])
        if when is None or when.date() != day:
            continue
        try:
            total += float(row[COL_KCAL - 1] or 0)
        except (TypeError, ValueError):
            continue
    return total


class Ledger:
    def __init__(self, path: Path) -> None:
        self.path = path

    def _append_once(self, record: MealRecord) -> int:
        if self.path.exists():
            wb = load_workbook(self.path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "紀錄"
            ws.append(HEADERS)
            ws.freeze_panes = "A2"
            for i, width in enumerate(COLUMN_WIDTHS, start=1):
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

        total = round(_daily_total(ws, record.when.date()) + record.calories_kcal)

        r = ws.max_row + 1  # 新的一列加在最後
        ws.cell(row=r, column=1, value=record.when).number_format = "yyyy-mm-dd hh:mm:ss"
        _write_text(ws.cell(row=r, column=2), record.uploader)
        _write_text(ws.cell(row=r, column=3), record.description)
        _write_text(ws.cell(row=r, column=4), record.ingredients)
        ws.cell(row=r, column=5, value=record.weight_g)
        ws.cell(row=r, column=6, value=record.calories_kcal)
        ws.cell(row=r, column=7, value=total)
        _write_text(ws.cell(row=r, column=8), record.photo_path)

        # 先存成暫存檔，再整個換掉正式檔：寫到一半出錯，正式檔不會變成壞掉的半成品
        tmp = self.path.with_name(self.path.name + ".tmp")
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)  # Excel 正開著這個檔案時，這裡會出現 PermissionError
        finally:
            tmp.unlink(missing_ok=True)
        return total

    async def append_meal(self, record: MealRecord) -> int:
        """加一列，回傳當日已攝取熱量（含這一餐）。寫不進去會警告並重試。"""
        for attempt in range(EXCEL_RETRIES + 1):
            try:
                return await asyncio.to_thread(self._append_once, record)
            except OSError as exc:
                if attempt == EXCEL_RETRIES:
                    log.error("Excel 寫入失敗，已放棄：%s", exc)
                    raise LedgerBusyError(str(exc)) from None
                log.warning(
                    "⚠️ Excel 寫入失敗（檔案可能正被打開著，請先關閉）：%s —— %d 秒後重試（%d/%d）",
                    exc, EXCEL_RETRY_DELAY_SECONDS, attempt + 1, EXCEL_RETRIES,
                )
                await asyncio.sleep(EXCEL_RETRY_DELAY_SECONDS)
        raise AssertionError("unreachable")
