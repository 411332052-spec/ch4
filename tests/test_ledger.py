import asyncio
from datetime import datetime

import pytest
from openpyxl import load_workbook

from foodlog import ledger as ledger_module
from foodlog.ledger import HEADERS, Ledger, LedgerBusyError, MealRecord


def record(when, kcal=500, uploader="Penny", description="便當", ingredients="白飯", photo="photos\\a.jpg"):
    return MealRecord(when, uploader, description, ingredients, 300, kcal, photo)


def add(ledger, rec):
    return asyncio.run(ledger.append_meal(rec))


def rows(path):
    return list(load_workbook(path).active.iter_rows(values_only=True))


def test_creates_file_with_header_and_appends(tmp_path):
    led = Ledger(tmp_path / "log.xlsx")
    add(led, record(datetime(2026, 10, 20, 8, 0, 0), 400))
    add(led, record(datetime(2026, 10, 20, 12, 0, 0), 600))
    data = rows(led.path)
    assert list(data[0]) == HEADERS
    assert len(data) == 3, "第二筆是加在後面，不是覆蓋"
    assert data[1][5] == 400 and data[2][5] == 600


def test_daily_total_is_channel_wide_and_resets_each_day(tmp_path):
    led = Ledger(tmp_path / "log.xlsx")
    assert add(led, record(datetime(2026, 10, 20, 8, 0), 400, uploader="A")) == 400
    assert add(led, record(datetime(2026, 10, 20, 12, 0), 600, uploader="B")) == 1000
    assert add(led, record(datetime(2026, 10, 20, 23, 59), 0, uploader="A")) == 1000  # 不是食物：熱量 0
    assert add(led, record(datetime(2026, 10, 21, 0, 1), 250, uploader="A")) == 250  # 00:00 換日
    assert [r[6] for r in rows(led.path)[1:]] == [400, 1000, 1000, 250]


def test_total_survives_a_fresh_ledger_object(tmp_path):
    path = tmp_path / "log.xlsx"
    add(Ledger(path), record(datetime(2026, 10, 20, 8, 0), 400))
    assert add(Ledger(path), record(datetime(2026, 10, 20, 9, 0), 100)) == 500


@pytest.mark.parametrize("evil", ["=1+1", "+cmd", "-2+3", "@SUM(A1)", '=HYPERLINK("http://x","y")', "\t=1"])
def test_protection_4_formulas_become_text(tmp_path, evil):
    led = Ledger(tmp_path / "log.xlsx")
    add(led, record(datetime(2026, 10, 20, 8, 0), uploader=evil, description=evil, ingredients=evil, photo=evil))
    ws = load_workbook(led.path).active
    for col in (2, 3, 4, 8):
        cell = ws.cell(row=2, column=col)
        assert cell.data_type == "s", f"第 {col} 欄不能是公式"
        assert cell.value.startswith("'")


def test_protection_4_leaves_normal_text_alone(tmp_path):
    led = Ledger(tmp_path / "log.xlsx")
    add(led, record(datetime(2026, 10, 20, 8, 0), uploader="Penny", description="雞腿便當"))
    ws = load_workbook(led.path).active
    assert ws["B2"].value == "Penny" and ws["C2"].value == "雞腿便當"


def test_newlines_collapse_to_one_line(tmp_path):
    led = Ledger(tmp_path / "log.xlsx")
    add(led, record(datetime(2026, 10, 20, 8, 0), description="a\nb\r\nc"))
    assert "\n" not in load_workbook(led.path).active["C2"].value


def test_excel_open_retries_then_gives_up_without_losing_old_data(tmp_path, monkeypatch, caplog):
    led = Ledger(tmp_path / "log.xlsx")
    add(led, record(datetime(2026, 10, 20, 8, 0), 400))
    monkeypatch.setattr(ledger_module, "EXCEL_RETRY_DELAY_SECONDS", 0)
    attempts = []

    def locked(src, dst):
        attempts.append(1)
        raise PermissionError("file is open in Excel")

    monkeypatch.setattr(ledger_module.os, "replace", locked)
    with pytest.raises(LedgerBusyError):
        add(led, record(datetime(2026, 10, 20, 9, 0), 100))
    assert len(attempts) == ledger_module.EXCEL_RETRIES + 1
    assert "⚠️" in caplog.text
    assert len(rows(led.path)) == 2, "舊資料不能壞"
    assert not (tmp_path / "log.xlsx.tmp").exists()


def test_excel_recovers_when_unlocked_midway(tmp_path, monkeypatch):
    led = Ledger(tmp_path / "log.xlsx")
    monkeypatch.setattr(ledger_module, "EXCEL_RETRY_DELAY_SECONDS", 0)
    real_replace = ledger_module.os.replace
    state = {"n": 0}

    def flaky(src, dst):
        state["n"] += 1
        if state["n"] < 3:
            raise PermissionError("busy")
        return real_replace(src, dst)

    monkeypatch.setattr(ledger_module.os, "replace", flaky)
    assert add(led, record(datetime(2026, 10, 20, 8, 0), 400)) == 400
    assert len(rows(led.path)) == 2
