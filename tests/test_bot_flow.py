"""不連 Discord，用假的訊息跑完整流程：收到 → ⏳ → 排隊 → 分析 → 存照片 → 建檔 → 回覆。"""
import asyncio
import logging
from dataclasses import replace
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from foodlog import bot as bot_module
from foodlog import replies
from foodlog.analysis import Analysis, AnalysisInvalid
from foodlog.bot import WAITING_EMOJI, FoodBot
from foodlog.ledger import LedgerBusyError

PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 40
CHANNEL = 1
FOOD = Analysis(True, "雞腿便當", ("白飯", "炸雞腿"), 650, 920)
MUG = Analysis(False, "馬克杯", (), 0, 0)


class FakeAttachment:
    def __init__(self, data=PNG, filename="a.png", content_type="image/png", size=None):
        self.data, self.filename, self.content_type = data, filename, content_type
        self.size = len(data) if size is None else size

    async def read(self):
        return self.data


class FakeMessage:
    _next_id = 100

    def __init__(self, attachments, user_id=7, name="Penny", channel=CHANNEL, bot=False, content=""):
        FakeMessage._next_id += 1
        self.id = FakeMessage._next_id
        self.attachments = attachments
        self.content = content
        self.channel = SimpleNamespace(id=channel)
        self.author = SimpleNamespace(id=user_id, display_name=name, bot=bot, mention=f"<@{user_id}>")
        self.created_at = datetime(2026, 10, 20, 6, 30, 5, tzinfo=timezone.utc)
        self.events: list[tuple] = []

    async def reply(self, text, **kwargs):
        self.events.append(("reply", text, kwargs))

    async def add_reaction(self, emoji):
        self.events.append(("add", emoji))

    async def remove_reaction(self, emoji, user):
        self.events.append(("remove", emoji))

    def replies(self):
        return [e[1] for e in self.events if e[0] == "reply"]

    def kinds(self):
        return [e[0] for e in self.events]


@pytest.fixture
def bot(cfg, monkeypatch):
    results: list = []

    async def fake_analyze(cfg_, data, ext, caption=""):
        result = results.pop(0)
        if isinstance(result, Exception):
            raise result
        return result

    monkeypatch.setattr(bot_module, "analyze_photo", fake_analyze)
    b = FoodBot(cfg)
    b.fake_results = results
    return b


def deliver(bot, *messages):
    async def main():
        bot.jobs.start()
        for m in messages:
            await bot.on_message(m)
        await bot.jobs._queue.join()
        await bot.jobs.stop()

    asyncio.run(main())


def rows(cfg):
    return list(load_workbook(cfg.excel_path).active.iter_rows(values_only=True))[1:]


def photo_names(cfg):
    return sorted(p.name for p in cfg.photos_dir.iterdir())


def test_food_photo_end_to_end(bot, cfg):
    bot.fake_results.append(FOOD)
    msg = FakeMessage([FakeAttachment()])
    deliver(bot, msg)

    assert msg.kinds() == ["add", "reply", "remove"], "先加 ⏳，回覆結果，最後拿掉 ⏳"
    assert ("add", WAITING_EMOJI) in msg.events
    text = msg.replies()[0]
    assert "雞腿便當" in text and "650" in text and "920" in text and "當日已攝取熱量" in text
    kwargs = [e for e in msg.events if e[0] == "reply"][0][2]
    assert kwargs["allowed_mentions"].everyone is False

    (row,) = rows(cfg)
    assert row[1] == "Penny" and row[2] == "雞腿便當" and row[3] == "白飯、炸雞腿"
    assert (row[4], row[5], row[6]) == (650, 920, 920)
    assert len(photo_names(cfg)) == 1
    assert photo_names(cfg)[0].endswith("雞腿便當.png") and row[7].endswith(photo_names(cfg)[0])


def test_not_food_is_still_recorded_saved_and_answered(bot, cfg):
    bot.fake_results.append(MUG)
    msg = FakeMessage([FakeAttachment()])
    deliver(bot, msg)

    assert "你傳的不是食物，是馬克杯" in msg.replies()[0]
    (row,) = rows(cfg)
    assert (row[2], row[3], row[4], row[5], row[6]) == ("不是食物", "馬克杯", 0, 0, 0)
    assert photo_names(cfg)[0].endswith("不是食物.png")


def test_daily_total_adds_up_across_people_and_photos(bot, cfg):
    bot.fake_results += [FOOD, MUG, FOOD]
    m1 = FakeMessage([FakeAttachment()], user_id=1, name="A")
    m2 = FakeMessage([FakeAttachment()], user_id=2, name="B")
    m3 = FakeMessage([FakeAttachment()], user_id=2, name="B")
    deliver(bot, m1, m2, m3)
    assert [r[6] for r in rows(cfg)] == [920, 920, 1840]
    assert "1,840" in m3.replies()[0]


def test_two_photos_in_one_message_are_two_meals(bot, cfg):
    bot.fake_results += [FOOD, FOOD]
    msg = FakeMessage([FakeAttachment(), FakeAttachment()])
    deliver(bot, msg)
    assert len(rows(cfg)) == 2 and len(msg.replies()) == 2
    assert msg.kinds().count("add") == 1 and msg.kinds().count("remove") == 1
    assert len(photo_names(cfg)) == 2


@pytest.mark.parametrize(
    "msg",
    [
        FakeMessage([]),
        FakeMessage([FakeAttachment(filename="a.pdf", content_type="application/pdf")]),
        FakeMessage([FakeAttachment()], channel=999),
        FakeMessage([FakeAttachment()], bot=True),
    ],
    ids=["text-only", "pdf", "other-channel", "bot-author"],
)
def test_ignored_messages_get_no_reaction_no_reply_no_record(bot, cfg, msg):
    deliver(bot, msg)
    assert msg.events == [] and not cfg.excel_path.exists() and photo_names(cfg) == []


def test_bad_analysis_replies_fail_message_and_records_nothing(bot, cfg):
    bot.fake_results.append(AnalysisInvalid("not json"))
    msg = FakeMessage([FakeAttachment()])
    deliver(bot, msg)
    assert msg.replies() == [f"<@7> {replies.FAIL_ANALYSIS}"]
    assert not cfg.excel_path.exists() and photo_names(cfg) == []
    assert msg.kinds().count("remove") == 1


def test_one_bad_photo_does_not_stop_the_next(bot, cfg):
    bot.fake_results += [RuntimeError("kaboom"), FOOD]
    m1, m2 = FakeMessage([FakeAttachment()]), FakeMessage([FakeAttachment()])
    deliver(bot, m1, m2)
    assert replies.FAIL_GENERIC in m1.replies()[0]
    assert "雞腿便當" in m2.replies()[0] and len(rows(cfg)) == 1


def test_fake_image_and_oversized_image_are_rejected_before_analysis(bot, cfg):
    fake = FakeMessage([FakeAttachment(data=b"%PDF-1.7 not an image", content_type="image/png")])
    big = FakeMessage([FakeAttachment(size=999_999_999)])
    deliver(bot, fake, big)
    assert replies.FAIL_BAD_IMAGE in fake.replies()[0]
    assert replies.FAIL_TOO_BIG in big.replies()[0]
    assert not cfg.excel_path.exists()


def test_excel_busy_removes_photo_and_tells_user(bot, cfg, monkeypatch):
    async def busy(record):
        raise LedgerBusyError("open in Excel")

    monkeypatch.setattr(bot.ledger, "append_meal", busy)
    bot.fake_results.append(FOOD)
    msg = FakeMessage([FakeAttachment()])
    deliver(bot, msg)
    assert replies.FAIL_LEDGER in msg.replies()[0]
    assert photo_names(cfg) == [], "沒建成檔就不留孤兒照片"


def test_too_many_from_one_person_gets_busy_reply(bot, cfg):
    bot.fake_results += [FOOD] * 3
    msg = FakeMessage([FakeAttachment() for _ in range(4)])
    deliver(bot, msg)
    assert len(rows(cfg)) == 3
    assert any(replies.BUSY_USER in r for r in msg.replies())


def log_lines(caplog):
    return [r.getMessage() for r in caplog.records if r.name.startswith("foodlog")]


def test_every_step_is_logged_in_order_with_one_tag(bot, cfg, caplog):
    caplog.set_level(logging.INFO)
    bot.fake_results.append(FOOD)
    msg = FakeMessage([FakeAttachment()])
    deliver(bot, msg)

    lines = log_lines(caplog)
    steps = ["收到", "已加入排隊", "輪到了", "已下載照片", "呼叫 claude -p", "分析結果", "照片已存", "已寫入 Excel", "已回覆上傳者", "處理完成"]
    positions = [next(i for i, line in enumerate(lines) if step in line) for step in steps]
    assert positions == sorted(positions), f"log 順序不對：{lines}"
    tagged = [line for line in lines if line.startswith(f"[{msg.id} 第1/1張]")]
    assert len(tagged) == len(steps) - 1, "除了「收到」，每一步都要帶同一個標籤"
    assert any("雞腿便當" in line and "920 大卡" in line for line in lines)
    assert any("當日累計 920 大卡" in line for line in lines)


def test_failures_are_logged_as_warnings(bot, cfg, caplog):
    caplog.set_level(logging.INFO)
    bot.fake_results.append(AnalysisInvalid("缺少欄位 name"))
    deliver(bot, FakeMessage([FakeAttachment()]))
    warnings = [r for r in caplog.records if r.levelno == logging.WARNING]
    assert any("分析失敗" in r.getMessage() and "缺少欄位 name" in r.getMessage() for r in warnings)


def test_rejections_are_logged(bot, cfg, caplog):
    caplog.set_level(logging.INFO)
    bot.fake_results += [FOOD] * 3
    deliver(bot, FakeMessage([FakeAttachment() for _ in range(4)]))
    assert any("沒有加入排隊" in line for line in log_lines(caplog))


def test_log_cannot_be_forged_by_newlines_in_names(bot, cfg, caplog):
    caplog.set_level(logging.INFO)
    bot.fake_results.append(Analysis(True, "牛排\n2026-01-01 ERROR 假的 log", ("肉",), 300, 700))
    deliver(bot, FakeMessage([FakeAttachment()], name="Evil\r\nINFO 假冒"))
    for line in log_lines(caplog):
        assert "\n" not in line and "\r" not in line


def test_token_never_appears_in_logs(bot, cfg, caplog):
    caplog.set_level(logging.DEBUG)
    bot.cfg = replace(cfg, discord_token="SECRET-TOKEN-123")
    bot.fake_results.append(FOOD)
    deliver(bot, FakeMessage([FakeAttachment()]))
    assert caplog.records, "測試本身要有東西可以檢查"
    assert "SECRET-TOKEN-123" not in caplog.text
